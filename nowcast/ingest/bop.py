"""Ingest Balance of Payments data from the CBU's BOP analytical presentation (Excel).

The file has quarterly data from 2005-Q1, in million USD.  We extract:

* ``bop_goods_exports``  row 6 "Goods, credits (exports)" — total goods exports
* ``bop_goods_imports``  row 7 "Goods, debits (imports)" — total goods imports
* ``bop_secondary_income_credits``  row 14 "Secondary income, credits" — remittances proxy

These become three catalogue series that feed the quarterly panel variables
``bop_exports``, ``bop_imports``, and extend ``remittances_real``.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pandas as pd

# Row numbers (1-indexed) for the three series in the "Dataset" sheet
SERIES_ROWS = {
    "bop_goods_exports": 6,
    "bop_goods_imports": 7,
    "bop_secondary_income_credits": 14,
}

HEADER_ROW = 4  # quarter labels like "2005-Q1"


def parse_bop_excel(path: str | Path) -> dict[str, pd.Series]:
    """Parse the BOP analytical presentation Excel into quarterly series.

    Returns a dict mapping series_id to a pd.Series with PeriodIndex (freq='Q').
    Values are in million USD.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["Dataset"]

    # Read quarter column positions from header row
    quarters: list[tuple[int, str]] = []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, c).value
        if v is not None:
            quarters.append((c, str(v)))

    result = {}
    for series_id, row_num in SERIES_ROWS.items():
        values = []
        periods = []
        for col_idx, q_label in quarters:
            val = ws.cell(row_num, col_idx).value
            if val is not None:
                # Convert "2005-Q1" to "2005Q1" for PeriodIndex
                period_str = q_label.replace("-", "")
                periods.append(period_str)
                values.append(float(val))
        idx = pd.PeriodIndex(periods, freq="Q")
        result[series_id] = pd.Series(values, index=idx, name=series_id)

    wb.close()
    return result


def bop_to_observations(
    path: str | Path,
    vintage_date: dt.date | None = None,
    run_id: str = "bop-seed",
) -> pd.DataFrame:
    """Convert BOP Excel into an observations DataFrame compatible with the vintage store.

    Each series gets item='total' and one row per quarter.
    """
    if vintage_date is None:
        vintage_date = dt.date(2026, 10, 19)  # date file was obtained

    series_dict = parse_bop_excel(path)
    rows = []
    now = dt.datetime.now(dt.timezone.utc)
    for series_id, s in series_dict.items():
        for period, value in s.items():
            rows.append({
                "series_id": series_id,
                "item": "total",
                "ref_period": str(period),
                "value": value,
                "vintage_date": vintage_date,
                "source_hash": "",
                "fetched_at": now,
                "run_id": run_id,
                "removed": False,
            })
    return pd.DataFrame(rows)


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "BOP_Analytical_Uzbekistan_1.xlsx"
    obs = bop_to_observations(path)
    print(f"Parsed {len(obs)} observations from BOP file")
    for sid in obs["series_id"].unique():
        sub = obs[obs["series_id"] == sid]
        print(f"  {sid}: {len(sub)} quarters, {sub['ref_period'].iloc[0]} to {sub['ref_period'].iloc[-1]}")
